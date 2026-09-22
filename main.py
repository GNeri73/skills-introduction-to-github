from pathlib import Path
import sys

from openpyxl import load_workbook


def main() -> int:
    file_name = input("Inserisci il nome del file Excel: ").strip()

    if not file_name:
        print("Nome file non valido.")
        return 1

    try:
        workbook = load_workbook(filename=Path(file_name), read_only=True, data_only=True)
    except FileNotFoundError:
        print("File non trovato.")
        return 1
    except Exception as exc:
        print(f"Errore durante l'apertura del file: {exc}")
        return 1

    try:
        if not workbook.worksheets:
            print("Il file Excel non contiene fogli leggibili.")
            return 1

        sheet = workbook.worksheets[0]
        print(sheet.cell(row=1, column=1).value)
        return 0
    except Exception as exc:
        print(f"Errore durante la lettura del file: {exc}")
        return 1
    finally:
        workbook.close()


if __name__ == "__main__":
    sys.exit(main())
